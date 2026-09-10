from pathlib import Path
import csv, hashlib, json, re
import numpy as np
import pandas as pd
from openpyxl import load_workbook

N=Path(__file__).resolve().parents[1]
P=N.parent
A=N/'02_acquired'
O=N/'03_analysis/PXD011817'
O.mkdir(exist_ok=True)
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def sheet(w,name):
    rows=list(w[name].values)
    return pd.DataFrame(rows[1:],columns=rows[0]).dropna(how='all')
w=load_workbook(A/'PXD011817_QUANT.xlsx',read_only=True,data_only=True)
raw=sheet(w,'proteinGroups')
elab=sheet(w,'proteinGroups_elaborated')
suppw=load_workbook(A/'Alfano2019_mmc5.xlsx',read_only=True,data_only=True)
supp=sheet(suppw,suppw.sheetnames[0])
lfq=[c for c in elab.columns if str(c).startswith('LFQ intensity Exp')]
assert len(lfq)==30 and len(supp)==105
src_cols=list(supp.columns[5:])
assert len(src_cols)==30
mapping=[]
for col,sc in zip(lfq,src_cols):
    m=re.search(r'Testicular ECM (\d+)_(\d+)',col)
    a=re.search(r'(Positive|Negative) Retrieval (\d+)_(\d+)',sc)
    assert m and a and m.groups()==a.groups()[1:]
    mapping.append(dict(column=col,source_header=sc,participant='ECM'+m[1],ecm_id=int(m[1]),technical_repeat=int(m[2]),group='SRpos' if a[1]=='Positive' else 'SRneg'))
indexed=elab.set_index('Protein IDs')
assert indexed.index.is_unique
checks=[]
for _,r in supp.iterrows():
    source_id=r['Protein IDs']
    cleaned_id=';'.join(x for x in source_id.split(';') if not x.startswith('CON__'))
    matched_id=source_id if source_id in indexed.index else cleaned_id
    v=indexed.loc[matched_id,lfq].to_numpy(dtype=float)
    s=r[src_cols].to_numpy(dtype=float)
    same=np.array_equal(v,s,equal_nan=True)
    checks.append(dict(protein_ids=source_id,PRIDE_elaborated_ids=matched_id,id_rule='exact' if matched_id==source_id else 'CON_accessions_removed_in_deposited_elaborated_sheet',lfq_values=30,all_exactly_match=bool(same)))
assert all(r['all_exactly_match'] for r in checks), "Supplementary LFQ does not match deposit"
pd.DataFrame(checks).to_csv(O/'supplement_to_PRIDE_value_check.tsv',sep='\t',index=False)
pd.DataFrame(mapping).to_csv(O/'technical_run_to_participant.tsv',sep='\t',index=False)
samples=pd.DataFrame([dict(participant='ECM'+str(i),ecm_id=i,group='SRpos' if i<=5 else 'SRneg',histology_source='heterogeneous residual spermatogenesis' if i<=5 else 'SCO') for i in range(1,11)])
samples.to_csv(O/'participants.tsv',sep='\t',index=False)

h=pd.read_csv(P/'01_raw/hgnc_complete_set.txt',sep='\t',dtype=str,keep_default_na=False)
h=h.loc[h.status.eq('Approved')]
approved=set(h.symbol)
aliases={}
for _,r in h.iterrows():
    for field in ['alias_symbol','prev_symbol']:
        for name in r[field].split('|'):
            if name: aliases.setdefault(name,set()).add(r.symbol)
def resolve_names(names):
    names=str(names or '')
    parts=[v.strip() for v in names.split(';') if v.strip()]
    if not parts:return None,'no_gene_annotation'
    resolved=[]
    for name in parts:
        choices={name} if name in approved else aliases.get(name,set())
        if len(choices)!=1:return None,'unmapped_or_ambiguous_gene_name'
        resolved.append(next(iter(choices)))
    if len(set(resolved))!=1:return None,'multiple_genes'
    return resolved[0],'single_gene_annotation'

rawid=raw.set_index('id')
assert rawid.index.is_unique
audit=[];values=[]
for _,r in elab.iterrows():
    rr=rawid.loc[r['id']]
    raw_ids=rr['Protein Ids']
    base_ids=lambda ids: {v.removeprefix('CON__').removeprefix('REV__') for v in ids.split(';')}
    assert base_ids(r['Protein IDs']) <= base_ids(raw_ids)
    raw_lfq=['LFQ intensity '+re.search(r'(Exp\d+)_',c)[1] for c in lfq]
    assert np.array_equal(np.array(r[lfq],dtype=float),np.array(rr[raw_lfq],dtype=float),equal_nan=True)
    gene,reason=resolve_names(r['Gene names'])
    peptides=float(rr['Peptides'] or 0); unique=float(rr['Unique peptides'] or 0)
    flags={c:str(rr.get(c,'') or '') for c in ['Reverse','Contaminant','Only identified by site']}
    has_flagged_accession=any(x.startswith(('CON__','REV__')) for x in raw_ids.split(';'))
    quality=(peptides>=2 and unique>=1 and not any(v=='+' for v in flags.values()) and not has_flagged_accession)
    v=np.array(r[lfq],dtype=float)
    v[v<=0]=np.nan
    v=np.log2(v).reshape(10,3)
    valid=(np.isfinite(v).sum(axis=1)>=2)
    agg=np.array([np.nanmedian(x) if good else np.nan for x,good in zip(v,valid)])
    row=dict(protein_ids=r['Protein IDs'],raw_protein_ids=raw_ids,source_id=r['id'],source_gene_names=r['Gene names'],gene=gene or '',mapping_status=reason,peptides=peptides,unique_peptides=unique,quality_pass=bool(quality),SRpos_n=int(np.isfinite(agg[:5]).sum()),SRneg_n=int(np.isfinite(agg[5:]).sum()),all_10_measured=bool(np.isfinite(agg).all()),median_log2_LFQ=float(np.nanmedian(agg)) if np.isfinite(agg).any() else None,selected_gene_group=False,**flags)
    audit.append(row);values.append(agg)
ad=pd.DataFrame(audit)
qualified=ad.loc[ad.quality_pass & ad.gene.ne('') & ad.median_log2_LFQ.notna()]
selected=qualified.sort_values(['gene','median_log2_LFQ','protein_ids'],ascending=[True,False,True]).drop_duplicates('gene').index
ad.loc[selected,'selected_gene_group']=True
ad.to_csv(O/'all_protein_group_annotation_audit.tsv',sep='\t',index=False)
X=pd.DataFrame(np.vstack(values)[selected],index=ad.loc[selected,'gene'],columns=samples.participant)
X.index.name='gene'
eligible=(X.iloc[:,:5].notna().sum(axis=1)>=3)&(X.iloc[:,5:].notna().sum(axis=1)>=3)
X=X.loc[eligible]
X.to_csv(O/'gene_log2_LFQ_min3_each_group.tsv',sep='\t',na_rep='NA')
complete=X.dropna()
complete.to_csv(O/'gene_log2_LFQ_complete_10.tsv',sep='\t')
sets=json.loads((P/'10_external_validation_20260910/01_metadata/fixed_gene_sets.json').read_text())
if 'gene_sets' in sets:sets=sets['gene_sets']
if 'primary' in sets:sets=sets['primary']
coverage=[]
for name,genes in sets.items():
    if isinstance(genes,dict):genes=genes.get('genes',genes.get('members'))
    members=sorted(set(genes)&set(complete.index))
    coverage.append(dict(program=name,fixed_members=len(genes),min3_each_group=len(set(genes)&set(X.index)),complete_case_members=len(members),genes=';'.join(members),status='EVALUABLE' if len(members)>=10 else 'NOT_EVALUABLE_COVERAGE'))
pd.DataFrame(coverage).to_csv(O/'fixed_program_coverage.tsv',sep='\t',index=False)
pd.DataFrame(dict(participant=X.columns,qualified_gene_measurements=X.notna().sum(axis=0),median_log2_LFQ=X.median(axis=0))).to_csv(O/'participant_coverage_QC.tsv',sep='\t',index=False)
summary=dict(status='PREPARED_AFTER_EXACT_SUPPLEMENT_MAPPING',comparison='SRneg minus SRpos within iNOA',participants=10,group_n={'SRpos':5,'SRneg':5},technical_replicates_per_patient=3,matched_supplement_protein_groups=len(checks),matched_numeric_values=len(checks)*30,all_values_exact=True,original_elaborated_groups=len(elab),gene_annotated_quality_groups=len(qualified),selected_gene_groups=len(selected),min3_each_group_genes=len(X),complete_case_genes=len(complete),coverage=coverage,source_sha256={p.name:digest(p) for p in [A/'PXD011817_QUANT.xlsx',A/'Alfano2019_mmc5.xlsx',N/'00_admin/PXD011817_ANALYSIS_PLAN.md']},limitations=['No individual clinical covariates or verified acquisition randomization in downloaded files.','Single-gene annotation excludes unresolved names but does not independently verify each accession.','Isolated ECM retrieval strata are not OA versus NOA.'])
(O/'preparation_receipt.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
print(json.dumps(summary,ensure_ascii=False,indent=2))
